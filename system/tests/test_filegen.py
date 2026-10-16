"""
Tests for bank file generation.

What is worth protecting here is not that a file gets written. It is that the
file says the same thing in every tab, that a value no bank would accept is
caught before the file leaves, that one merchant's missing data does not hold
back the other fifty, and that a template asking for something the code cannot
produce fails loudly instead of writing a blank.
"""

from datetime import date

import pytest
from openpyxl import load_workbook

from config.acquirers import ACQUIRERS
from config.templates import TEMPLATES
from src.catalogue.eligibility import check
from src.catalogue.mcc import CATALOGUE, Instrument
from src.filegen.cover import draft
from src.filegen.fill import (
    DERIVATIONS,
    BatchMeta,
    Context,
    fill,
    instrument_set,
)
from src.filegen.preflight import Severity, blocking_rows, preflight
from src.filegen.spec import Column, Grain, Necessity, Source, format_date
from src.filegen.write import write
from src.identifiers import gstin_checksum
from src.model import (
    Merchant,
    PriorityTier,
    ProcurementRequest,
    Segment,
    Stream,
)

CC = Instrument.CREDIT_CARD
DC = Instrument.DEBIT_CARD
EMI = Instrument.EMI

_GST_BODY = "29AABCA1234K1Z"
GOOD_GST = _GST_BODY + gstin_checksum(_GST_BODY)

DEAL = {
    "settlement_cycle": "Daily",
    "expected_transaction": "5000000",
    "relationship_manager": "Priya Mehta",
    "business_model": "B2C",
    "point_of_contact": "Merchant",
    "sales_hod": "Vikram Iyer",
    "zonal_head": "Arjun Reddy",
}

META = BatchMeta(lot_no="20261001-TEST-1", login_date=date(2026, 10, 1))


def merchant(merchant_id: str = "M-00001", name: str = "Ambervale Wellness Private Limited") -> Merchant:
    return Merchant(
        merchant_id=merchant_id,
        legal_name=name,
        segment=Segment.SME,
        priority_tier=PriorityTier.P3,
        website="https://ambervale.test",
        fields={
            "entity_legal_name": name,
            "company_pan": "AABCA1234K",
            "company_gst": GOOD_GST,
            "registered_address": "14 Mill Road, Bengaluru 560001",
            "registration_number": "U51909KA2019PTC123456",
            "signatory_name": "Meera Iyer",
            "signatory_pan": "ABCPI9876L",
            "signatory_address": "8 Church Street, Bengaluru 560001",
            "website": "https://ambervale.test",
            "state": "Karnataka",
            "state_code": "29",
            "city": "Bengaluru",
            # Leading zero on purpose: the value a spreadsheet silently eats.
            "pincode": "045001",
            "entity_type": "Private Limited",
            "incorporation_date": "2019-04-03",
            "contact_name": "Meera Iyer",
            "contact_mobile": "9812345678",
            "contact_email": "accounts@ambervale.test",
            "settlement_bank": "Thornfield Bank",
            "settlement_account_name": name,
            "settlement_account_number": "156974319638",
            "settlement_ifsc": "THNF0123456",
        },
    )


def request(
    acquirer: str = "NORTHWAY_CYBS",
    mcc: str = "7101",
    instruments: tuple = (CC, DC),
    m: Merchant | None = None,
) -> ProcurementRequest:
    m = m or merchant()
    return ProcurementRequest(
        request_id=f"REQ-00001-{mcc}-{acquirer}",
        merchant=m,
        mcc=mcc,
        acquirer=ACQUIRERS[acquirer],
        stream=Stream.BAU,
        instruments=instruments,
    )


def build(acquirer: str = "NORTHWAY_CYBS", requests=None, deals=None):
    requests = requests or [request(acquirer)]
    eligibility = {r.request_id: check(acquirer, r.mcc, r.instruments) for r in requests}
    deals = deals if deals is not None else {
        r.merchant.merchant_id: dict(DEAL) for r in requests
    }
    return fill(acquirer, requests, META, eligibility, deals=deals)


# --- the templates themselves -------------------------------------------


def test_every_derived_column_has_a_rule():
    """A template is config and the rules are code, so the two can drift. This
    is the check that stops a new column quietly writing a blank into a
    mandatory cell and the file bouncing at the bank a week later."""
    for code, template in TEMPLATES.items():
        for sheet in template.sheets:
            for column in sheet.columns:
                if column.source in (Source.DERIVED, Source.OPS_INPUT) and column.key:
                    assert column.key in DERIVATIONS, (
                        f"{code}/{sheet.name}: {column.header!r} wants "
                        f"{column.key!r}, which has no rule"
                    )


def test_every_dropdown_column_names_a_domain_the_template_defines():
    for code, template in TEMPLATES.items():
        for sheet in template.sheets:
            for column in sheet.columns:
                if column.domain:
                    assert column.domain in template.domains, (
                        f"{code}: {column.header!r} is validated against "
                        f"{column.domain!r}, which the template does not define"
                    )


def test_no_template_is_blocked_on_a_column_nobody_owns():
    """`mandatory_gaps` is the honest answer to how much of a template can be
    generated. A rate card column marked mandatory would mean no file ever goes
    out without sales typing in it, and that has to be a deliberate decision
    rather than a config accident."""
    for code, template in TEMPLATES.items():
        unowned = [
            c for c in template.mandatory_gaps
            if c.source not in (Source.SALES_OWNED, Source.FINANCE_OWNED)
        ]
        assert not unowned, f"{code}: {[c.header for c in unowned]} belong to nobody"


def test_the_templates_disagree_with_each_other():
    """The dataset exists to model that every bank's form is different. If the
    templates converge, the mapping problem disappears and the phase proves
    nothing."""
    formats = {t.date_format for t in TEMPLATES.values()}
    assert len(formats) >= 3
    assert {t.file_format for t in TEMPLATES.values()} == {"xlsx", "csv"}
    assert len({len(t.sheets) for t in TEMPLATES.values()}) > 1


def test_one_bank_on_two_platforms_has_two_templates():
    cybs, fss = TEMPLATES["NORTHWAY_CYBS"], TEMPLATES["NORTHWAY_FSS"]
    assert [s.name for s in cybs.sheets] != [s.name for s in fss.sheets]
    assert cybs.date_format != fss.date_format


# --- mapping ------------------------------------------------------------


def test_the_mcc_comes_from_the_request_not_the_merchant():
    """The MCC belongs to the use case. Read off the merchant record it is
    blank, and it is blank in a mandatory column on every row."""
    f = build("CASTERLY", [request("CASTERLY", mcc="7407")])
    assert f.rows[0].values["mcc"] == "7407"


def test_one_acquirer_gets_the_state_code_and_another_the_state_name():
    cybs = build("NORTHWAY_CYBS").rows[0].values
    fss = build("NORTHWAY_FSS").rows[0].values
    assert cybs["State"] == "Karnataka"
    assert fss["State Code"] == "29"


def test_each_acquirer_gets_its_own_date_format():
    aldridge = build("ALDRIDGE").rows[0].values["Business Registration Date"]
    fss = build("NORTHWAY_FSS").rows[0].values["Business Registration Date"]
    assert aldridge == "03/04/2019"
    assert fss == "2019-04-03"


def test_a_date_is_never_left_ambiguous_between_formats():
    """03-04-2026 is 3 April to one bank and 4 March to another, and neither
    file looks wrong."""
    d = date(2026, 4, 3)
    assert format_date(d, "%d-%m-%Y") == "03-04-2026"
    assert format_date(d, "%Y-%m-%d") == "2026-04-03"
    assert format_date(d, "%d-%b-%Y") == "03-Apr-2026"


def test_instrument_sets_are_written_in_a_fixed_order():
    """The same set must not produce two strings, or a dropdown match becomes
    luck."""
    assert instrument_set((DC, CC)) == "CC,DC"
    assert instrument_set((CC, DC)) == "CC,DC"
    assert instrument_set((CC, DC, EMI)) == "CC,DC,EMI"


def test_the_file_states_what_the_acquirer_will_actually_run():
    """A conditional verdict can close one rail and leave the others open.
    Sending only the requested set gets the whole row refused."""
    f = build("BRINDLE", [request("BRINDLE", mcc="7726", instruments=(CC, DC))])
    row = f.rows[0].values
    assert row["Instruments Requested"] == "CC,DC"
    assert row["Instruments Permitted"] == "DC"


def test_an_industry_is_mapped_to_the_forms_own_vocabulary():
    f = build("ALDRIDGE", [request("ALDRIDGE", mcc="7101")])
    assert f.rows[0].values["Industry"] == "Retail"


def test_nothing_is_invented_for_a_cell_another_team_owns():
    """An empty settlement cycle stays empty. A plausible value put in by ops
    is a commercial term nobody agreed, and it goes live."""
    f = build("ALDRIDGE", [request("ALDRIDGE")], deals={})
    assert f.rows[0].values["Settlement Cycle"] == ""
    assert any(u.owner is Source.SALES_OWNED for u in f.unfilled)


def test_a_derived_column_with_no_rule_raises():
    from src.filegen.fill import _cell

    c = Context(
        request=request(),
        entry=CATALOGUE["7101"],
        eligibility=check("NORTHWAY_CYBS", "7101", (CC, DC)),
        meta=META,
        deal=dict(DEAL),
    )
    bad = Column("Invented", Source.DERIVED, Necessity.MANDATORY, key="no_such_rule")
    with pytest.raises(KeyError, match="no rule"):
        _cell(bad, c)


def test_an_unknown_condition_raises_rather_than_being_assumed_true():
    from src.filegen.fill import _applies

    c = Context(
        request=request(),
        entry=CATALOGUE["7101"],
        eligibility=check("NORTHWAY_CYBS", "7101", (CC, DC)),
        meta=META,
    )
    bad = Column("Invented", Source.OPS_INPUT, Necessity.CONDITIONAL,
                 key="licence_ref", applies_when="no_such_condition")
    with pytest.raises(KeyError, match="CONDITIONS"):
        _applies(bad, c)


def test_a_conditional_column_is_only_reported_where_it_applies():
    """A licence reference is needed by the merchants whose category requires a
    licence, not by all fifty. Reported on all fifty, it buries the findings
    that matter."""
    clean = build("ALDRIDGE", [request("ALDRIDGE", mcc="7101")])
    licensed = build("ALDRIDGE", [request("ALDRIDGE", mcc="7407")])
    assert not [u for u in clean.unfilled if u.header == "Licence Reference"]
    assert [u for u in licensed.unfilled if u.header == "Licence Reference"]


# --- sheet grain and internal consistency --------------------------------


def test_a_merchant_with_two_use_cases_appears_once_on_the_identity_tab():
    """Twice on the identity tab asks the bank to open two accounts for one
    merchant, and what comes back is two customer ids to unpick later."""
    m = merchant()
    requests = [
        request("NORTHWAY_CYBS", mcc="7101", m=m),
        request("NORTHWAY_CYBS", mcc="7104", m=m),
    ]
    requests[1].request_id = "REQ-00001-7104-NORTHWAY_CYBS"
    f = build("NORTHWAY_CYBS", requests)
    assert len(f.rows_for("Accounts")) == 1
    assert len(f.rows_for("Deals")) == 2


def test_grain_is_declared_on_every_sheet_that_repeats_the_identity_block():
    cybs = TEMPLATES["NORTHWAY_CYBS"]
    for name in ("Accounts", "Discrepancy", "Contacts"):
        sheet = next(s for s in cybs.sheets if s.name == name)
        assert sheet.grain is Grain.MERCHANT
    deals = next(s for s in cybs.sheets if s.name == "Deals")
    assert deals.grain is Grain.APPLICATION


def test_the_two_identity_tabs_cannot_disagree():
    """The error the hand-filled version makes most often: the same merchant
    described two ways in two tabs of one file."""
    f = build("NORTHWAY_CYBS")
    result = preflight(f)
    assert not [x for x in result.findings if x.check == "tabs disagree"]
    assert not [x for x in result.findings if x.check == "mirror gap"]


def test_a_tab_that_disagrees_is_caught():
    f = build("NORTHWAY_CYBS")
    disc = f.rows_for("Discrepancy")[0]
    disc.values["Firm PAN No"] = "ZZZZZ9999Z"
    result = preflight(f)
    assert [x for x in result.findings if x.check == "tabs disagree"]
    assert not result.can_send


# --- pre-flight ---------------------------------------------------------


def test_a_value_one_word_off_the_dropdown_is_a_blocker():
    """The check that pays for the module. The bank's upload compares the cell
    against its list exactly, and nothing about the file looks wrong."""
    deal = dict(DEAL, settlement_cycle="T+1")
    f = build("BRINDLE", [request("BRINDLE")], deals={"M-00001": deal})
    result = preflight(f)
    off = [x for x in result.blockers if x.check == "off-list value"]
    assert off
    assert "T+1" in off[0].message


def test_the_message_names_the_value_that_was_probably_meant():
    deal = dict(DEAL, settlement_cycle="daily")
    f = build("BRINDLE", [request("BRINDLE")], deals={"M-00001": deal})
    off = [x for x in preflight(f).blockers if x.check == "off-list value"]
    assert "'Daily'" in off[0].message


def test_a_loose_dropdown_is_a_warning_not_a_stop():
    """One acquirer's upload takes any value and queries it later. That is a
    slower, cheaper failure than a refused file, and it deserves a different
    answer."""
    deal = dict(DEAL, point_of_contact="Partner")
    f = build("NORTHWAY_CYBS", [request("NORTHWAY_CYBS")], deals={"M-00001": deal})
    result = preflight(f)
    off = [x for x in result.findings if x.check == "off-list value"]
    assert off and off[0].severity is Severity.WARNING
    assert result.can_send


def test_an_empty_mandatory_cell_stops_the_file():
    f = build("ALDRIDGE", [request("ALDRIDGE")], deals={})
    assert not preflight(f).can_send


def test_the_row_cap_is_split_not_truncated():
    """A file one row over the cap is refused in full, and the quiet version of
    this bug drops the overflow with nobody knowing which merchants went."""
    cap = TEMPLATES["BRINDLE"].max_rows
    requests = []
    for i in range(cap + 1):
        m = merchant(f"M-{i:05d}", f"Harrowgate {i} Private Limited")
        r = request("BRINDLE", m=m)
        r.request_id = f"REQ-{i:05d}-7101-BRINDLE"
        requests.append(r)
    f = build("BRINDLE", requests)
    cap_findings = [x for x in preflight(f).blockers if x.check == "row cap"]
    assert cap_findings
    assert "Split into 2 files" in cap_findings[0].message


def test_a_line_break_in_a_csv_value_is_a_blocker():
    """One address with a line break shifts every column after it."""
    m = merchant()
    m.fields["registered_address"] = "14 Mill Road\nBengaluru 560001"
    f = build("CASTERLY", [request("CASTERLY", m=m)])
    assert [x for x in preflight(f).blockers if x.check == "csv safety"]


def test_numbers_a_spreadsheet_would_rewrite_are_flagged():
    f = build("NORTHWAY_CYBS")
    flags = [x for x in preflight(f).findings if x.check == "number formatting"]
    assert flags
    assert all(x.severity is Severity.WARNING for x in flags)


def test_a_duplicate_application_is_caught():
    f = build("CASTERLY")
    f.rows.append(f.rows[0])
    assert [x for x in preflight(f).blockers if x.check == "duplicate row"]


def test_a_clean_batch_passes():
    f = build("VERRAN")
    result = preflight(f)
    assert result.can_send
    assert result.summary() in ("Clean.", "1 to look at; send.")


# --- holding rows back rather than the file ------------------------------


def test_one_merchants_missing_data_does_not_hold_the_others():
    """Same reasoning as partial raising at intake: the next slot is tomorrow,
    and the clock on the good rows keeps running."""
    good = merchant("M-00001", "Harrowgate Trading Private Limited")
    bad = merchant("M-00002", "Verdant Supplies Private Limited")
    requests = [request("ALDRIDGE", m=good), request("ALDRIDGE", m=bad)]
    requests[1].request_id = "REQ-00002-7101-ALDRIDGE"

    f = build("ALDRIDGE", requests, deals={"M-00001": dict(DEAL)})
    dropped = blocking_rows(f)
    assert list(dropped) == ["REQ-00002-7101-ALDRIDGE"]
    assert "sales" in dropped["REQ-00002-7101-ALDRIDGE"]


def test_a_file_level_problem_is_not_blamed_on_a_row():
    """A row cap is a problem with the file. Dropping merchants would not fix
    it, so it must not appear as a row-level block."""
    cap = TEMPLATES["BRINDLE"].max_rows
    requests = []
    for i in range(cap + 1):
        m = merchant(f"M-{i:05d}", f"Harrowgate {i} Private Limited")
        r = request("BRINDLE", m=m)
        r.request_id = f"REQ-{i:05d}-7101-BRINDLE"
        requests.append(r)
    f = build("BRINDLE", requests)
    assert blocking_rows(f) == {}


# --- the written file ---------------------------------------------------


def test_the_workbook_has_every_tab_and_a_dropdown_sheet(tmp_path):
    f = build("NORTHWAY_CYBS")
    path = write(f, tmp_path, "LOT1", "2026-10-01")
    wb = load_workbook(path)
    assert [s.name for s in f.template.sheets] == wb.sheetnames[:-1]
    assert wb.sheetnames[-1] == "Drop down list"


def test_identifiers_survive_the_spreadsheet(tmp_path):
    """A pincode starting with zero and a twelve-digit account number are both
    numbers to a spreadsheet. Stored as numbers, the bank reads something else
    than what was reviewed."""
    f = build("NORTHWAY_CYBS")
    path = write(f, tmp_path, "LOT1", "2026-10-01")
    ws = load_workbook(path)["Accounts"]
    headers = [c.value for c in ws[1]]
    cell = ws.cell(row=2, column=headers.index("Zipcode") + 1)
    assert cell.value == "045001"
    assert isinstance(cell.value, str)
    assert cell.number_format == "@"


def test_the_dropdowns_are_real_validation(tmp_path):
    f = build("NORTHWAY_CYBS")
    path = write(f, tmp_path, "LOT1", "2026-10-01")
    ws = load_workbook(path)["Accounts"]
    assert ws.data_validations.dataValidation


def test_the_csv_holds_one_row_per_application(tmp_path):
    import csv as csvlib

    m = merchant()
    requests = [request("CASTERLY", mcc="7101", m=m), request("CASTERLY", mcc="7104", m=m)]
    requests[1].request_id = "REQ-00001-7104-CASTERLY"
    f = build("CASTERLY", requests)
    path = write(f, tmp_path, "LOT1", "2026-10-01")
    rows = list(csvlib.DictReader(path.open()))
    assert len(rows) == 2
    assert {r["mcc"] for r in rows} == {"7101", "7104"}


def test_the_file_name_carries_the_acquirer_and_the_lot(tmp_path):
    f = build("VERRAN")
    path = write(f, tmp_path, "20261001-VERR-1", "2026-10-01")
    assert path.name == "2026-10-01_VERRAN_lot20261001-VERR-1.xlsx"


# --- the covering email -------------------------------------------------


def test_the_draft_names_the_conditional_merchants():
    """A condition left in a column is read by nobody, and the merchant goes
    live on a rail the policy does not allow."""
    r = request("ALDRIDGE", mcc="7407")
    f = build("ALDRIDGE", [r])
    eligibility = {r.request_id: check("ALDRIDGE", "7407", (CC, DC))}
    note = draft(
        f, preflight(f), attachment="x.xlsx", lot_no="LOT1", on=date(2026, 10, 1),
        eligibility=eligibility, acquirer_name="Aldridge Bank",
        acquirer_email="aldridge.ops@example-acquirer.test",
    )
    assert "conditional category" in note.body
    assert "Ambervale Wellness Private Limited" in note.body


def test_the_internal_note_is_marked_for_deletion():
    f = build("NORTHWAY_CYBS")
    checked = preflight(f)
    note = draft(
        f, checked, attachment="x.xlsx", lot_no="LOT1", on=date(2026, 10, 1),
        eligibility={}, acquirer_name="Northway Bank",
        acquirer_email="northway.ops@example-acquirer.test",
    )
    if checked.warnings:
        assert "delete before sending" in note.body
    assert note.subject.startswith("MID Request | Lot LOT1")


def test_nothing_in_this_module_sends_anything():
    """The one deliberate checkpoint at the external boundary. A draft is a
    file on disk; a person sends it."""
    import inspect

    from src.filegen import cover, write as writer

    for module in (cover, writer):
        source = inspect.getsource(module)
        for forbidden in ("smtplib", "sendmail", "requests.post", "urlopen"):
            assert forbidden not in source
